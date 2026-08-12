from merchant_map_cache import build_merchant_cache_key

REVIEW_QUEUE_SHEET_NAME = "Review_Queue"
REVIEW_QUEUE_HEADERS = [
    "Merchant_Cache_Key",
    "Sample_Narration",
    "Occurrences",
    "Total_Debit",
    "Total_Credit",
    "Sources",
    "Suggested_Category",
    "Suggested_Subcategory",
    "Resolved_Category",
    "Resolved_Subcategory",
]
RESOLVED_CATEGORY_COLUMN_INDEX = 9
RESOLVED_SUBCATEGORY_COLUMN_INDEX = 10


def build_account_source_label(source_account):
    return "card" if source_account.startswith("icici_credit_card") else "bank"


def aggregate_review_candidates(transactions):
    review_candidates = {}

    for transaction in transactions:
        if transaction["is_duplicate"] or not transaction["needs_review"]:
            continue

        merchant_cache_key = build_merchant_cache_key(transaction)
        candidate = review_candidates.setdefault(
            merchant_cache_key,
            {
                "sample_narration": transaction["narration_clean"],
                "occurrences": 0,
                "total_debit": 0.0,
                "total_credit": 0.0,
                "sources": set(),
                "suggested_category": transaction["category"],
                "suggested_subcategory": transaction["subcategory"],
            },
        )

        candidate["occurrences"] += 1
        candidate["sources"].add(build_account_source_label(transaction["source_account"]))
        amount = transaction["amount"]
        if amount is not None:
            if transaction["direction"] == "debit":
                candidate["total_debit"] += float(amount)
            else:
                candidate["total_credit"] += float(amount)

    return review_candidates


def build_review_queue_rows(transactions, already_queued_merchant_keys=()):
    review_candidates = aggregate_review_candidates(transactions)
    already_queued = set(already_queued_merchant_keys)

    review_queue_rows = []
    for merchant_cache_key, candidate in review_candidates.items():
        if merchant_cache_key in already_queued:
            continue
        review_queue_rows.append(
            [
                merchant_cache_key,
                candidate["sample_narration"],
                candidate["occurrences"],
                candidate["total_debit"] or None,
                candidate["total_credit"] or None,
                "/".join(sorted(candidate["sources"])),
                candidate["suggested_category"],
                candidate["suggested_subcategory"],
                None,
                None,
            ]
        )

    review_queue_rows.sort(key=lambda row: (-(row[3] or 0) - (row[4] or 0), row[0]))
    return review_queue_rows


def read_queued_merchant_keys(workbook_file_path):
    from pathlib import Path

    from openpyxl import load_workbook

    workbook_path = Path(workbook_file_path)
    if not workbook_path.exists():
        return set()

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if REVIEW_QUEUE_SHEET_NAME not in workbook.sheetnames:
            return set()
        worksheet = workbook[REVIEW_QUEUE_SHEET_NAME]
        queued_keys = set()
        for row in worksheet.iter_rows(min_row=2, min_col=1, max_col=1, values_only=True):
            if row[0]:
                queued_keys.add(str(row[0]).strip())
        return queued_keys
    finally:
        workbook.close()


def read_resolved_review_entries(workbook_file_path, category_taxonomy):
    from pathlib import Path

    from openpyxl import load_workbook

    from category_taxonomy import is_valid_category_pair

    workbook_path = Path(workbook_file_path)
    if not workbook_path.exists():
        return {}, []

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if REVIEW_QUEUE_SHEET_NAME not in workbook.sheetnames:
            return {}, []

        worksheet = workbook[REVIEW_QUEUE_SHEET_NAME]
        resolved_classifications = {}
        rejected_resolutions = []

        for row in worksheet.iter_rows(min_row=2, min_col=1, max_col=len(REVIEW_QUEUE_HEADERS), values_only=True):
            merchant_cache_key = row[0]
            if not merchant_cache_key:
                continue

            resolved_category = row[RESOLVED_CATEGORY_COLUMN_INDEX - 1]
            resolved_subcategory = row[RESOLVED_SUBCATEGORY_COLUMN_INDEX - 1]
            if not resolved_category or not resolved_subcategory:
                continue

            resolved_category = str(resolved_category).strip()
            resolved_subcategory = str(resolved_subcategory).strip()

            if not is_valid_category_pair(category_taxonomy, resolved_category, resolved_subcategory):
                rejected_resolutions.append(
                    {
                        "merchant_cache_key": str(merchant_cache_key).strip(),
                        "category": resolved_category,
                        "subcategory": resolved_subcategory,
                        "reason": "Not a valid category and subcategory pair in the taxonomy",
                    }
                )
                continue

            resolved_classifications[str(merchant_cache_key).strip()] = {
                "category": resolved_category,
                "subcategory": resolved_subcategory,
                "confidence": "high",
                "reason": "Resolved by hand in the review queue",
            }

        return resolved_classifications, rejected_resolutions
    finally:
        workbook.close()


def append_review_queue_rows(workbook_file_path, rows):
    if not rows:
        return None

    for row in rows:
        if len(row) != len(REVIEW_QUEUE_HEADERS):
            raise ValueError(f"Review row has {len(row)} values, expected {len(REVIEW_QUEUE_HEADERS)}")

    from pathlib import Path

    import xlwings

    workbook_path = Path(workbook_file_path).resolve()
    excel_was_already_running = bool(xlwings.apps)
    workbook_was_already_open = any(
        Path(open_book.fullname).resolve() == workbook_path
        for excel_app in xlwings.apps
        for open_book in excel_app.books
    )

    workbook = xlwings.Book(str(workbook_path))
    try:
        if REVIEW_QUEUE_SHEET_NAME in [sheet.name for sheet in workbook.sheets]:
            worksheet = workbook.sheets[REVIEW_QUEUE_SHEET_NAME]
        else:
            worksheet = workbook.sheets.add(REVIEW_QUEUE_SHEET_NAME, after=workbook.sheets[-1])
            worksheet.range((1, 1)).value = REVIEW_QUEUE_HEADERS

        last_populated_row = worksheet.range((worksheet.cells.last_cell.row, 1)).end("up").row
        first_empty_row = max(last_populated_row + 1, 2)
        worksheet.range((first_empty_row, 1)).value = rows
        workbook.save()
        return first_empty_row, first_empty_row + len(rows) - 1
    finally:
        if not workbook_was_already_open:
            excel_application = workbook.app
            workbook.close()
            if not excel_was_already_running and not excel_application.books:
                excel_application.quit()