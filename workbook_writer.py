from pathlib import Path

UNCATEGORIZED_CATEGORY = "Uncategorized"
TRANSACTIONS_SHEET_NAME = "Detailed_Transactions"
TRANSACTION_IDENTIFIER_COLUMN_INDEX = 11
TRANSACTION_DATE_COLUMN_INDEX = 1
TRANSACTION_DATE_NUMBER_FORMAT = "DD-MM-YYYY"
ACCOUNT_DISPLAY_NAMES = {
    "icici_savings": "ICICI Savings",
    "icici_credit_card_amazon": "ICICI Amazon Card",
    "icici_credit_card_standard": "ICICI Credit Card",
}
WORKBOOK_TRANSACTION_HEADERS = [
    "Transaction_Date",
    "Account_Type",
    "Transaction_Narration",
    "Debit_Amount",
    "Credit_Amount",
    "Month_Year",
    "Assigned_Category",
    "Specific_Expense_Type",
    "Classification_Confidence",
    "Classification_Reason",
    "Transaction_ID",
]


def build_month_year(transaction_date):
    if transaction_date is None:
        return None
    return transaction_date.strftime("%Y-%m")


def build_account_display_name(source_account):
    return ACCOUNT_DISPLAY_NAMES.get(source_account, source_account)


def build_workbook_row(transaction):
    amount = transaction["amount"]
    debit_amount = float(amount) if amount is not None and transaction["direction"] == "debit" else None
    credit_amount = float(amount) if amount is not None and transaction["direction"] == "credit" else None

    return [
        transaction["transaction_date"],
        build_account_display_name(transaction["source_account"]),
        transaction["narration_clean"],
        debit_amount,
        credit_amount,
        build_month_year(transaction["transaction_date"]),
        transaction["category"],
        transaction["subcategory"],
        transaction.get("classification_confidence"),
        transaction.get("classification_reason"),
        transaction["transaction_id"],
    ]


def build_workbook_rows(transactions):
    return [build_workbook_row(transaction) for transaction in transactions if not transaction["is_duplicate"]]


def read_written_transaction_identifiers(workbook_file_path):
    from openpyxl import load_workbook

    workbook_path = Path(workbook_file_path)
    if not workbook_path.exists():
        return set()

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if TRANSACTIONS_SHEET_NAME not in workbook.sheetnames:
            raise ValueError(f"Sheet {TRANSACTIONS_SHEET_NAME!r} not found in {workbook_path.name}")
        worksheet = workbook[TRANSACTIONS_SHEET_NAME]
        written_identifiers = set()
        for row in worksheet.iter_rows(
            min_row=2,
            min_col=TRANSACTION_IDENTIFIER_COLUMN_INDEX,
            max_col=TRANSACTION_IDENTIFIER_COLUMN_INDEX,
            values_only=True,
        ):
            identifier = row[0]
            if identifier:
                written_identifiers.add(str(identifier).strip())
        return written_identifiers
    finally:
        workbook.close()


def summarize_workbook_rows(rows):
    months = sorted({row[5] for row in rows if row[5]})
    accounts = sorted({row[1] for row in rows})
    total_debit = sum(row[3] for row in rows if row[3])
    total_credit = sum(row[4] for row in rows if row[4])
    uncategorized_count = sum(1 for row in rows if row[6] == "Uncategorized")
    transfer_count = sum(1 for row in rows if row[6] == "Transfers")

    return {
        "row_count": len(rows),
        "months": months,
        "accounts": accounts,
        "total_debit": total_debit,
        "total_credit": total_credit,
        "uncategorized_count": uncategorized_count,
        "transfer_count": transfer_count,
    }


def format_workbook_summary(summary):
    lines = [
        f"rows to append      : {summary['row_count']}",
        f"months              : {', '.join(summary['months'])}",
        f"accounts            : {', '.join(summary['accounts'])}",
        f"total debit         : {summary['total_debit']:,.2f}",
        f"total credit        : {summary['total_credit']:,.2f}",
        f"transfer rows       : {summary['transfer_count']}",
        f"uncategorized rows  : {summary['uncategorized_count']}",
    ]
    return "\n".join(lines)


def verify_workbook_headers(worksheet):
    header_values = worksheet.range((1, 1), (1, len(WORKBOOK_TRANSACTION_HEADERS))).value
    actual_headers = [str(value).strip() if value is not None else "" for value in header_values]
    if actual_headers != WORKBOOK_TRANSACTION_HEADERS:
        raise ValueError(
            f"Workbook headers do not match expected layout.\n"
            f"expected: {WORKBOOK_TRANSACTION_HEADERS}\n"
            f"found   : {actual_headers}"
        )


def find_first_empty_row(worksheet):
    last_populated_row = worksheet.range((worksheet.cells.last_cell.row, 1)).end("up").row
    return max(last_populated_row + 1, 2)


def append_transaction_rows(workbook_file_path, rows):
    if not rows:
        return None

    for row in rows:
        if len(row) != len(WORKBOOK_TRANSACTION_HEADERS):
            raise ValueError(f"Row has {len(row)} values, expected {len(WORKBOOK_TRANSACTION_HEADERS)}")

    import xlwings

    workbook_path = Path(workbook_file_path).resolve()
    excel_was_already_running = bool(xlwings.apps)
    workbook_was_already_open = any(
        Path(open_book.fullname).resolve() == workbook_path
        for excel_app in xlwings.apps
        for open_book in excel_app.books
    )

    workbook = xlwings.Book(str(workbook_path))
    try:
        worksheet = workbook.sheets[TRANSACTIONS_SHEET_NAME]
        verify_workbook_headers(worksheet)
        first_empty_row = find_first_empty_row(worksheet)
        last_written_row = first_empty_row + len(rows) - 1

        worksheet.range((first_empty_row, 1)).value = rows
        worksheet.range(
            (first_empty_row, TRANSACTION_DATE_COLUMN_INDEX),
            (last_written_row, TRANSACTION_DATE_COLUMN_INDEX),
        ).number_format = TRANSACTION_DATE_NUMBER_FORMAT
        workbook.save()
        return first_empty_row, last_written_row
    finally:
        if not workbook_was_already_open:
            excel_application = workbook.app
            workbook.close()
            if not excel_was_already_running and not excel_application.books:
                excel_application.quit()


CATEGORY_COLUMN_INDEX = 7
CLASSIFICATION_REASON_COLUMN_INDEX = 10


def build_category_update_block(existing_category_rows, transactions):
    transactions_by_identifier = {
        transaction["transaction_id"]: transaction for transaction in transactions
    }

    updated_block = []
    changed_row_summaries = []

    for row_offset, existing_row in enumerate(existing_category_rows):
        category, subcategory, confidence, reason, transaction_identifier = existing_row
        transaction = transactions_by_identifier.get(transaction_identifier)

        if transaction is None:
            updated_block.append([category, subcategory, confidence, reason])
            continue

        if transaction["category"] == UNCATEGORIZED_CATEGORY and category != UNCATEGORIZED_CATEGORY:
            updated_block.append([category, subcategory, confidence, reason])
            continue

        new_values = [
            transaction["category"],
            transaction["subcategory"],
            transaction.get("classification_confidence"),
            transaction.get("classification_reason"),
        ]
        updated_block.append(new_values)

        if [category, subcategory] != new_values[:2]:
            changed_row_summaries.append(
                {
                    "row_number": row_offset + 2,
                    "transaction_id": transaction_identifier,
                    "from_category": f"{category} > {subcategory}",
                    "to_category": f"{transaction['category']} > {transaction['subcategory']}",
                }
            )

    return updated_block, changed_row_summaries


def read_existing_category_rows(workbook_file_path):
    from openpyxl import load_workbook

    workbook_path = Path(workbook_file_path)
    if not workbook_path.exists():
        return []

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        worksheet = workbook[TRANSACTIONS_SHEET_NAME]
        existing_rows = []
        for row in worksheet.iter_rows(
            min_row=2,
            min_col=CATEGORY_COLUMN_INDEX,
            max_col=TRANSACTION_IDENTIFIER_COLUMN_INDEX,
            values_only=True,
        ):
            if row[-1]:
                existing_rows.append(list(row))
        return existing_rows
    finally:
        workbook.close()


def apply_category_updates(workbook_file_path, updated_block):
    if not updated_block:
        return 0

    import xlwings

    workbook_path = Path(workbook_file_path).resolve()
    excel_was_already_running = bool(xlwings.apps)
    workbook_was_already_open = any(
        Path(open_book.fullname).resolve() == workbook_path
        for excel_app in xlwings.apps
        for open_book in excel_app.books
    )

    workbook = xlwings.Book(str(workbook_path))
    try:
        worksheet = workbook.sheets[TRANSACTIONS_SHEET_NAME]
        verify_workbook_headers(worksheet)
        last_row = len(updated_block) + 1
        worksheet.range(
            (2, CATEGORY_COLUMN_INDEX),
            (last_row, CLASSIFICATION_REASON_COLUMN_INDEX),
        ).value = updated_block
        workbook.save()
        return len(updated_block)
    finally:
        if not workbook_was_already_open:
            excel_application = workbook.app
            workbook.close()
            if not excel_was_already_running and not excel_application.books:
                excel_application.quit()